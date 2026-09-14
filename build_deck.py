#!/usr/bin/env python3
"""
Build the Field CTO X-Arch executive deck (.pptx) from the slide outline.

16:9 dark-theme deck: title + 12 condensed core slides (each with a tailored
visual) + appendix divider + 8 appendix slides. Bullets are tightened to short
fragments; full speaker notes are embedded in each slide's notes pane. All
figures are illustrative (see 06-financial-model.md).

Run:  python3 build_deck.py
Out:  Field-CTO-XArch-Business-Case.pptx
"""

import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# ---- Theme ----------------------------------------------------------------
BG = RGBColor(0x10, 0x10, 0x14)
PANEL = RGBColor(0x1A, 0x1A, 0x2E)
CHIP = RGBColor(0x24, 0x24, 0x3C)
ACCENT = RGBColor(0x00, 0xBC, 0xEB)
ACCENT2 = RGBColor(0x00, 0xA4, 0xFD)
TEXT = RGBColor(0xF2, 0xF4, 0xF8)
MUTED = RGBColor(0x9A, 0xA3, 0xB2)
LINE = RGBColor(0x2C, 0x2C, 0x3A)
FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


# ---- Primitives -----------------------------------------------------------
def new_slide():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    r.fill.solid(); r.fill.fore_color.rgb = BG; r.line.fill.background()
    r.shadow.inherit = False
    sp = r._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2, sp)
    return s


def add_text(slide, left, top, width, height, runs, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, space_after=6):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after)
        for (text, size, color, bold) in para:
            run = p.add_run(); run.text = text
            run.font.size = Pt(size); run.font.color.rgb = color
            run.font.bold = bold; run.font.name = FONT
    return tb


def shp(slide, shape, l, t, w, h, fill, line_color=None, line_w=0.75):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color; sp.line.width = Pt(line_w)
    return sp


def stext(sp, text, size, color, bold=True, align=PP_ALIGN.CENTER):
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.color.rgb = color
    run.font.bold = bold; run.font.name = FONT


def conn(slide, x1, y1, x2, y2, color=ACCENT, w=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                   Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w); c.shadow.inherit = False
    return c


def panel_bg(slide, l, t, w, h):
    return shp(slide, MSO_SHAPE.RECTANGLE, l, t, w, h, PANEL, line_color=LINE, line_w=1)


def bullets_para(items, size, color=TEXT):
    return [[(u"\u2022  ", size, ACCENT, True), (it, size, color, False)] for it in items]


def content_slide(kicker, headline, bullets, notes, visual=None, footer=None):
    s = new_slide()
    add_text(s, Inches(0.6), Inches(0.45), Inches(12.1), Inches(0.4),
             [[(kicker.upper(), 12, ACCENT, True)]])
    shp(s, MSO_SHAPE.RECTANGLE, 0.6, 0.95, 1.4, 0.055, ACCENT)
    add_text(s, Inches(0.6), Inches(1.15), Inches(12.1), Inches(1.1),
             [[(headline, 27, TEXT, True)]])
    shp(s, MSO_SHAPE.RECTANGLE, 0.6, 2.37, 12.13, 0.014, LINE)
    if visual is not None:
        panel_bg(s, 0.6, 2.55, 6.0, 4.3)
        add_text(s, Inches(0.9), Inches(2.78), Inches(5.4), Inches(3.9),
                 bullets_para(bullets, 16), space_after=11)
        panel_bg(s, 6.85, 2.55, 5.85, 4.3)
        visual(s, 6.85, 2.55, 5.85, 4.3)
    else:
        panel_bg(s, 0.6, 2.55, 12.1, 4.3)
        add_text(s, Inches(1.0), Inches(2.8), Inches(11.3), Inches(3.9),
                 bullets_para(bullets, 18), space_after=12)
    shp(s, MSO_SHAPE.RECTANGLE, 0.6, 6.96, 12.13, 0.012, LINE)
    note = footer or "Illustrative figures \u2014 replace with actuals before presenting."
    add_text(s, Inches(0.6), Inches(7.03), Inches(10.6), Inches(0.32),
             [[("FIELD CTO \u00b7 X-ARCH SELLING", 9, MUTED, True),
               ("      " + note, 9, MUTED, False)]])
    s.notes_slide.notes_text_frame.text = notes
    return s


# ---- Visuals (each draws inside the right panel box l,t,w,h in inches) -----
def v_silos(s, l, t, w, h):
    names = ["Networking", "Security", "Observability", "SecOps"]
    gx, gy = l + 0.55, t + 0.35
    cw = (w - 1.1 - 0.3) / 2; ch = 0.62
    for i, nm in enumerate(names):
        x = gx + (i % 2) * (cw + 0.3); y = gy + (i // 2) * (ch + 0.2)
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch, CHIP, MUTED), nm, 12, TEXT)
    ay = gy + 2 * ch + 0.2 + 0.03
    stext(shp(s, MSO_SHAPE.DOWN_ARROW, l + w / 2 - 0.25, ay, 0.5, 0.42, ACCENT), "", 8, BG)
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, l + 0.55, ay + 0.55, w - 1.1, 0.8, ACCENT),
          "One integrated outcome", 15, BG)


def v_chips(items, cols=1, accents=frozenset()):
    def draw(s, l, t, w, h):
        n = len(items); rows = math.ceil(n / cols); pad = 0.4
        gx, gy = l + pad, t + 0.4; gw = w - 2 * pad; gh = h - 0.8
        cw = (gw - (cols - 1) * 0.3) / cols
        ch = (gh - (rows - 1) * 0.25) / rows
        for i, it in enumerate(items):
            x = gx + (i % cols) * (cw + 0.3); y = gy + (i // cols) * (ch + 0.25)
            fill = ACCENT if i in accents else CHIP
            stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch, fill,
                      None if i in accents else MUTED),
                  it, 13, BG if i in accents else TEXT, i in accents)
    return draw


def v_vflow(steps, size=12, accent_last=False, accent_all=False):
    def draw(s, l, t, w, h):
        n = len(steps); pad = 0.4; ah = 0.26
        top = t + 0.3; avail = h - 0.6
        bh = (avail - (n - 1) * ah) / n
        x = l + pad; bw = w - 2 * pad; cx = l + w / 2; y = top
        for i, st in enumerate(steps):
            hot = accent_all or (accent_last and i == n - 1)
            fill = ACCENT if hot else CHIP
            stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh, fill,
                      None if hot else MUTED), st, size, BG if hot else TEXT, hot)
            if i < n - 1:
                shp(s, MSO_SHAPE.DOWN_ARROW, cx - 0.16, y + bh + 0.02, 0.32, ah - 0.05, ACCENT)
            y += bh + ah
    return draw


def v_hub(s, l, t, w, h):
    cx, cy = l + w / 2, t + h / 2
    cw, ch = 1.7, 0.78
    spokes = ["CxO", "Networking", "Splunk", "Security", "Account Exec"]
    Rx, Ry, sw, sh = 1.95, 1.45, 1.45, 0.55
    for i, lbl in enumerate(spokes):
        rad = math.radians(-90 + i * (360 / len(spokes)))
        sx = cx + Rx * math.cos(rad); sy = cy + Ry * math.sin(rad)
        conn(s, cx, cy, sx, sy, ACCENT2, 1.25)
    for i, lbl in enumerate(spokes):
        rad = math.radians(-90 + i * (360 / len(spokes)))
        sx = cx + Rx * math.cos(rad) - sw / 2; sy = cy + Ry * math.sin(rad) - sh / 2
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, sw, sh, CHIP, MUTED), lbl, 11, TEXT, False)
    stext(shp(s, MSO_SHAPE.OVAL, cx - cw / 2, cy - ch / 2, cw, ch, ACCENT), "Field CTO", 14, BG)


def v_breakthrough(s, l, t, w, h):
    steps = [("Stalled / SI-locked whale", CHIP), ("BREAKTHROUGH \u00b7 ~60\u2013120 days", ACCENT),
             ("Durable multi-year growth", ACCENT2)]
    pad, ah = 0.4, 0.34; top = t + 0.35; avail = h - 0.85
    bh = (avail - 2 * ah) / 3; x = l + pad; bw = w - 2 * pad; cx = l + w / 2; y = top
    for i, (lbl, color) in enumerate(steps):
        cold = color == CHIP
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh, color, MUTED if cold else None),
              lbl, 13, TEXT if cold else BG, not cold)
        if i < 2:
            shp(s, MSO_SHAPE.DOWN_ARROW, cx - 0.18, y + bh + 0.02, 0.36, ah - 0.05, ACCENT)
        y += bh + ah
    add_text(s, Inches(l + 0.4), Inches(y + 0.02), Inches(w - 0.8), Inches(0.35),
             [[("Always hands the account back", 11, MUTED, False)]], align=PP_ALIGN.CENTER)


def v_focus(s, l, t, w, h):
    pad = 0.35; gx = l + pad; gw = w - 2 * pad
    add_text(s, Inches(gx), Inches(t + 0.18), Inches(gw), Inches(0.28),
             [[("3 + 1 ARCHITECTURES", 10, MUTED, True)]])
    pn = ["Networking", "Observe", "SecOps", "+1 AI"]
    cw = (gw - 3 * 0.2) / 4
    for i, p in enumerate(pn):
        hot = i < 3
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, gx + i * (cw + 0.2), t + 0.5, cw, 0.7,
                  ACCENT if hot else ACCENT2), p, 10.5, BG)
    add_text(s, Inches(gx), Inches(t + 1.4), Inches(gw), Inches(0.28),
             [[("4 VERTICALS", 10, MUTED, True)]])
    vn = ["Transportation", "Financial Svcs", "Healthcare", "Public Sector"]
    cw2 = (gw - 0.25) / 2
    for i, v in enumerate(vn):
        x = gx + (i % 2) * (cw2 + 0.25); y = t + 1.72 + (i // 2) * (0.7 + 0.22)
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw2, 0.7, CHIP, MUTED), v, 12, TEXT, False)


def v_stack(s, l, t, w, h):
    layers = [("OT \u2014 read-only / passive", ACCENT2),
              ("Secure Networking \u2014 segmentation + SD-WAN", ACCENT),
              ("Observability \u2014 Splunk OT / ITSI", ACCENT2),
              ("Resilience \u2014 Splunk ES + Cisco XDR", ACCENT)]
    n = len(layers); pad = 0.4; top = t + 0.35; avail = h - 0.7; gap = 0.18
    bh = (avail - (n - 1) * gap) / n; x = l + pad; bw = w - 2 * pad
    for i, (lbl, color) in enumerate(layers):
        y = top + (n - 1 - i) * (bh + gap)
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, bw, bh, color), lbl, 12, BG)


def v_bars(s, l, t, w, h):
    data = [("Conserv.", 9), ("Base", 31), ("Upside", 46)]
    maxv = 52; pad = 0.55
    add_text(s, Inches(l + 0.3), Inches(t + 0.12), Inches(w - 0.6), Inches(0.3),
             [[("Net profit by scenario ($M, illustrative)", 11, MUTED, False)]], align=PP_ALIGN.CENTER)
    base_y = t + h - 0.95; top_y = t + 0.55; plot_h = base_y - top_y
    n = len(data); slot = (w - 2 * pad) / n; bw = slot * 0.5
    for i, (lbl, val) in enumerate(data):
        bx = l + pad + i * slot + (slot - bw) / 2; bh = plot_h * (val / maxv); by = base_y - bh
        shp(s, MSO_SHAPE.RECTANGLE, bx, by, bw, bh, ACCENT if lbl == "Base" else ACCENT2)
        add_text(s, Inches(bx - 0.2), Inches(by - 0.34), Inches(bw + 0.4), Inches(0.3),
                 [[("$%dM" % val, 12, TEXT, True)]], align=PP_ALIGN.CENTER)
        add_text(s, Inches(bx - 0.25), Inches(base_y + 0.05), Inches(bw + 0.5), Inches(0.3),
                 [[(lbl, 11, MUTED, False)]], align=PP_ALIGN.CENTER)
    conn(s, l + pad - 0.1, base_y, l + w - pad + 0.1, base_y, MUTED, 1.0)
    add_text(s, Inches(l + 0.3), Inches(base_y + 0.42), Inches(w - 0.6), Inches(0.3),
             [[("Break-even \u2248 $9M bookings (~5 deals) \u2014 every scenario clears it", 9.5, MUTED, False)]],
             align=PP_ALIGN.CENTER)


def v_team(s, l, t, w, h):
    roles = ["8 \u00d7 Field CTO", "1 \u00d7 Group lead", "2 \u00d7 X-Arch architect", "1 \u00d7 Deal desk / ops"]
    pad = 0.4; gx = l + pad; gy = t + 0.4; gw = w - 2 * pad
    cw = (gw - 0.3) / 2; ch = 0.8
    for i, rl in enumerate(roles):
        x = gx + (i % 2) * (cw + 0.3); y = gy + (i // 2) * (ch + 0.25)
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, ch, CHIP, MUTED), rl, 12, TEXT, False)
    oy = gy + 2 * ch + 0.25 + 0.3
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, gx, oy, gw, 0.85, ACCENT),
          "~$5.4M / yr  \u00b7  Q3 go/expand/stop gate", 14, BG)


def v_gate(s, l, t, w, h):
    cx = l + w / 2; cy = t + 1.15
    stext(shp(s, MSO_SHAPE.DIAMOND, cx - 0.95, cy - 0.7, 1.9, 1.4, ACCENT), "Q3 GATE", 14, BG)
    cby = cy + 1.05; bw = w / 2 - 0.8
    conn(s, cx - 0.35, cy + 0.6, l + 0.5 + bw / 2, cby, ACCENT2, 1.25)
    conn(s, cx + 0.35, cy + 0.6, cx + 0.3 + bw / 2, cby, MUTED, 1.25)
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, l + 0.5, cby, bw, 0.7, ACCENT2),
          "Continue / expand", 12, BG)
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx + 0.3, cby, bw, 0.7, CHIP, MUTED),
          "Stop \u2014 spend capped", 12, TEXT, False)
    add_text(s, Inches(l + 0.5), Inches(cby + 0.95), Inches(w - 1.0), Inches(0.7),
             [[("Proof: touched vs. control \u2014 attach \u00b7 deal size \u00b7 win rate", 11, MUTED, False)]],
             align=PP_ALIGN.CENTER)


def v_window(s, l, t, w, h):
    pad = 0.4
    add_text(s, Inches(l + pad), Inches(t + 0.16), Inches(w - 2 * pad), Inches(0.28),
             [[("THE POST-SPLUNK WINDOW", 10, MUTED, True)]])
    tx = l + pad; tw = w - 2 * pad; ty = t + 1.2; th = 0.7; seg = tw / 4
    stext(shp(s, MSO_SHAPE.RECTANGLE, tx, ty, seg, th, CHIP, LINE), "Splunk close", 10, MUTED, False)
    stext(shp(s, MSO_SHAPE.RECTANGLE, tx + seg, ty, 2 * seg, th, ACCENT), "OPEN WINDOW \u2014 NOW", 12, BG)
    stext(shp(s, MSO_SHAPE.RECTANGLE, tx + 3 * seg, ty, seg, th, CHIP, LINE), "Commoditized", 10, MUTED, False)
    mx = tx + 2 * seg
    add_text(s, Inches(mx - 1.3), Inches(ty - 0.82), Inches(2.6), Inches(0.3),
             [[("First-mover advantage", 11, ACCENT, True)]], align=PP_ALIGN.CENTER)
    shp(s, MSO_SHAPE.DOWN_ARROW, mx - 0.15, ty - 0.44, 0.3, 0.4, ACCENT)
    add_text(s, Inches(tx), Inches(ty + th + 0.16), Inches(tw), Inches(0.28),
             [[("FORCES PULLING IT FORWARD", 10, MUTED, True)]])
    drivers = ["Portfolio", "Integrations", "Consolidate", "AI build-outs"]
    dy = ty + th + 0.46; dw = (tw - 3 * 0.2) / 4
    for i, d in enumerate(drivers):
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, tx + i * (dw + 0.2), dy, dw, 0.62, CHIP, MUTED),
              d, 11, TEXT, False)


def v_appendix(badge, srclabel):
    def draw(s, l, t, w, h):
        cx = l + w / 2; bw = 1.7; bh = 1.4
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - bw / 2, t + 0.55, bw, bh, CHIP, ACCENT, 1.25),
              badge, 40, ACCENT)
        add_text(s, Inches(l + 0.4), Inches(t + 2.2), Inches(w - 0.8), Inches(0.4),
                 [[("Reserve for Q&A", 15, MUTED, False)]], align=PP_ALIGN.CENTER)
        stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 2.0, t + 3.0, 4.0, 0.62, PANEL, LINE),
              "Source \u00b7 " + srclabel, 12, TEXT, False)
    return draw


def v_timeline(s, l, t, w, h):
    nodes = [("Approve", "Day 0"), ("+30d", "Lead \u00b7 charter \u00b7 comp"),
             ("+60d", "Cohort + accounts"), ("+90d", "First framings")]
    pad = 0.7; y = t + h / 2; x0 = l + pad; x1 = l + w - pad
    conn(s, x0, y, x1, y, ACCENT, 2.0)
    n = len(nodes)
    for i, (a, b) in enumerate(nodes):
        nx = x0 + (x1 - x0) * (i / (n - 1))
        shp(s, MSO_SHAPE.OVAL, nx - 0.13, y - 0.13, 0.26, 0.26, ACCENT)
        add_text(s, Inches(nx - 0.75), Inches(y - 0.8), Inches(1.5), Inches(0.45),
                 [[(a, 13, TEXT, True)]], align=PP_ALIGN.CENTER)
        add_text(s, Inches(nx - 0.9), Inches(y + 0.22), Inches(1.8), Inches(0.7),
                 [[(b, 10, MUTED, False)]], align=PP_ALIGN.CENTER)


# ---- Title slide ----------------------------------------------------------
s = new_slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.6), SLIDE_W, Inches(2.3))
band.fill.solid(); band.fill.fore_color.rgb = PANEL; band.line.fill.background()
band.shadow.inherit = False
add_text(s, Inches(0.8), Inches(2.75), Inches(11.7), Inches(0.5),
         [[("EXECUTIVE FUNDING DECISION", 14, ACCENT, True)]])
add_text(s, Inches(0.8), Inches(3.25), Inches(11.7), Inches(1.2),
         [[("Field CTO X-Arch Selling", 40, TEXT, True)]])
add_text(s, Inches(0.8), Inches(4.25), Inches(11.7), Inches(0.6),
         [[("Funding the motion that monetizes Cisco + Splunk", 20, MUTED, False)]])
_loop = ["Sense", "Connect", "Secure", "Understand", "Act"]
_cw, _gap = 1.85, 0.55
_tot = len(_loop) * _cw + (len(_loop) - 1) * _gap
_x0 = (13.333 - _tot) / 2
for _i, _st in enumerate(_loop):
    _x = _x0 + _i * (_cw + _gap)
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, _x, 5.35, _cw, 0.6, CHIP, ACCENT, 1.0), _st, 13, ACCENT)
    if _i < len(_loop) - 1:
        shp(s, MSO_SHAPE.RIGHT_ARROW, _x + _cw + 0.08, 5.35 + 0.19, _gap - 0.16, 0.22, ACCENT)
add_text(s, Inches(0.8), Inches(6.7), Inches(11.7), Inches(0.4),
         [[("Illustrative business case \u2014 figures are parameterized placeholders.", 11, MUTED, False)]])
s.notes_slide.notes_text_frame.text = (
    "Opening: We now own the most complete portfolio in the market. This deck asks "
    "leadership to fund a small, senior team whose only job is to sell it as one "
    "integrated outcome. It's a bounded, high-leverage bet with a Q3 stop gate. Keep "
    "the room focused on the decision: fund the pilot.")

# ---- Core slides (12) -----------------------------------------------------
content_slide(
    "The problem",
    "We own the most complete portfolio \u2014 and still sell it in silos",
    ["Four Cisco conversations, not one outcome",
     "Speeds-and-feeds, not funded outcomes",
     "Splunk + Cisco attach under-realized",
     "SIs own the 'architecture of the deal'",
     "Inaction: nine-figure pipeline, forgone yearly"],
    "The core tension: portfolio completeness without a motion to sell it as one. "
    "Land the line: 'the customer experiences four Cisco conversations, not one.' "
    "Then make inaction concrete: the GAP between a flat renewal and an X-Arch "
    "reframe, repeated across ~50 strategic accounts a year, is recurring, not "
    "one-time. Source: 01 \u00a71 (problem + cost-of-inaction box).",
    visual=v_silos)

content_slide(
    "Why now",
    "The post-Splunk window is open \u2014 and closing",
    ["One vendor, edge to SOC",
     "Integrations live and shipping",
     "Consolidation \u2192 a C-suite frame",
     "AI build-outs pull it all forward"],
    "Urgency without hype. The integrations are real and fresh; consolidation and "
    "AI builds create cross-architecture entry points right now. First mover on the "
    "motion compounds. Source: 01 \u00a72.",
    visual=v_window)

content_slide(
    "The differentiator",
    "One vendor, edge to SOC \u2014 if someone tells it end-to-end",
    ["Every arrow between pillars = a cross-sell",
     "Others own one box; we own the arrows",
     "Splunk + Cisco: top under-executed play",
     "The story rivals can't match end-to-end"],
    "The story competitors can't match end-to-end. Walk the loop once, slowly, and "
    "point at the arrows \u2014 that's where the incremental value lives. Source: 02 \u00a72.",
    visual=v_vflow(["Sense", "Connect", "Secure", "Understand", "Act"], accent_all=True))

content_slide(
    "The solution",
    "A thin, senior Field CTO seller group that leads across architectures",
    ["Owns the CxO outcome conversation",
     "Architects across BUs; specialists close",
     "Vertical-credible (NIS2/DORA/HIPAA)",
     "Orchestrates \u2014 not another overlay SE"],
    "Define the role crisply and pre-empt the 'another overlay team' objection: this "
    "is leverage that makes existing teams MORE productive. Thin, senior, "
    "orchestration-first. Source: 01 \u00a73; 03 \u00a71.",
    visual=v_hub)

content_slide(
    "Extending the model",
    "Black Ops: an elite tiger team for the accounts nobody can crack",
    ["Marathon coverage + breakthrough mission",
     "For stalled / SI-locked / narrow-window whales",
     "~60\u2013120 days \u2192 durable multi-year growth",
     "Always hands the account back",
     "Ethical; OT read-only guardrails hold"],
    "A mode, not separate headcount. Two things to stress: the mission is focused "
    "but the point is durable, multi-year growth for Cisco (it always hands the "
    "account back \u2014 no hoarding), and 'Black Ops' is a special-forces posture "
    "metaphor \u2014 fully ethical, nothing hidden from the customer. Source: 09.",
    visual=v_breakthrough)

content_slide(
    "Focus",
    "Three high-attach architectures, four high-pull verticals",
    ["3 anchor pillars + 1 adjacency",
     "Chosen for cross-attach + regulatory pull",
     "Joint Cisco + Splunk strength",
     "The list is swappable"],
    "We deliberately don't boil the ocean: three anchor pillars with the clearest "
    "boardroom narrative, plus AI-ready infra as adjacency. Four verticals chosen "
    "for high cross-attach, regulatory/resilience pull, and joint Cisco+Splunk "
    "strength \u2014 swappable without changing the operating design. Source: 02 \u00a74\u20135; "
    "04 master matrix.",
    visual=v_focus)

content_slide(
    "The motion",
    "A repeatable lifecycle and a library of plays \u2014 not heroics",
    ["A repeatable, six-stage lifecycle",
     "A library of packaged plays",
     "Play = script + architecture + value + proof",
     "Scales; survives key-person risk"],
    "Show this is systematized, not dependent on a few heroes. The play library "
    "codifies the knowledge so the motion scales and survives key-person risk. "
    "Source: 05 \u00a71\u20132.",
    visual=v_vflow(["Select", "Frame", "Architect", "Value", "Close", "Expand"], size=11, accent_last=True))

content_slide(
    "Worked example \u00b7 aviation",
    "Airport operator: Black Ops broke open a stalled 'smart airport' program",
    ["OT hub + NIS2 + a Terminal 4 build",
     "SI about to lock \u2014 pod convened",
     "Reframed as board-level resilience",
     "One architecture across the stack",
     "Won OT trust, then handed back"],
    "The end-to-end proof. Walk the transformation: stalled + SI-locked \u2192 executive "
    "breakthrough \u2192 one co-authored architecture \u2192 multi-year program. Stress the OT "
    "safety boundary as the trust-winner with airport engineering. Source: 10.",
    visual=v_stack)

content_slide(
    "The return",
    "~$180M influenced pipeline, ~$60M bookings, payback < 12 months",
    ["~$180M influenced pipeline (Yr 1)",
     "~$60M bookings \u2192 ~$36M margin",
     "Break-even ~$9M \u2014 about five deals",
     "Profitable across every scenario"],
    "The money slide. Anchor on break-even (~$9M bookings, about five deals, covers "
    "the entire $5.4M cost); everything above is upside. All three scenarios clear "
    "break-even \u2014 the question is pace, not payback. Illustrative and parameterized. "
    "Source: 01 \u00a76; 06 \u00a72\u20133, \u00a75.",
    visual=v_bars)

content_slide(
    "The ask",
    "Fund a 12-person, 4-quarter pilot (~$5.4M illustrative)",
    ["12 people, 4 quarters",
     "Internal-first staffing",
     "~40\u201360 accounts + paired controls",
     "Q3 go/expand/stop gate"],
    "State the ask plainly and immediately bound the downside: small, senior team; "
    "defined spend; a real stop gate at Q3. Ask for the decision. Source: 01 \u00a75; "
    "06 \u00a71.",
    visual=v_team)

content_slide(
    "Bounded risk + proof",
    "A Q3 gate and control-account measurement make this low-risk",
    ["Q3 gate caps the downside",
     "Control accounts prove the deltas",
     "Talent / comp / attribution mitigated",
     "Influence ledger prevents disputes"],
    "Address the skeptic. We measure with control accounts, so influence is proven "
    "not asserted; and we can stop at Q3 if the deltas aren't there. Source: 07 "
    "\u00a71\u20133; 05 \u00a77.",
    visual=v_gate)

content_slide(
    "Recommendation",
    "Approve the pilot; first executive framings within 90 days",
    ["Approve: 12 people, 4 quarters, Q3 gate",
     "+30d: lead, charter, comp, metrics",
     "+60d: cohort + accounts selected",
     "+90d: first framings; enablement on"],
    "Close with a clear call to action and a concrete near-term timeline so approval "
    "feels actionable, not open-ended. Restate: bounded downside, differentiated "
    "upside. Ask for the yes. Source: 01 \u00a79; README one-screen summary.",
    visual=v_timeline)

# ---- Appendix -------------------------------------------------------------
# (kicker, headline, bullets, notes, badge, source-label, TOC-label)
APPX = [
    ("Appendix \u00b7 A1", "Role, competencies, and skills matrix",
     ["4 pillars: business, X-arch, exec presence, orchestration",
      "Deep in \u22652 pillars; conversant in all",
      "Deep in one vertical; team covers four"],
     "Detail for HR/org-design questions. Source: 03 \u00a71\u20133.",
     "A1", "03 \u00a71\u20133", "A1 \u00b7 Role & skills"),
    ("Appendix \u00b7 A2", "Org placement and RACI",
     ["GTM-primary; dotted line to CTO/eng",
      "Revenue-accountable AND BU-neutral",
      "Field CTO owns frame; AE owns account"],
     "For 'where does it sit / who owns what' questions. Source: 03 \u00a75, \u00a78.",
     "A2", "03 \u00a75, \u00a78", "A2 \u00b7 Org & RACI"),
    ("Appendix \u00b7 A3", "Per-vertical field engagement kits",
     ["Buyers, discovery, hypotheses, objections",
      "One kit per vertical (x4)",
      "Strategy \u2192 field-ready playbook"],
     "For 'how do we actually sell this' questions. Source: 04 \u00a71\u20134.",
     "A3", "04 \u00a71\u20134", "A3 \u00b7 Vertical kits"),
    ("Appendix \u00b7 A4", "OT safety boundary",
     ["Read-only / passive acquisition only",
      "Never write to PLC/HMI/DCS/RTU/SIS",
      "Guardrail AND trust-winning differentiator"],
     "For OT/operations stakeholders. Source: 04 OT safety boundary.",
     "A4", "04 OT boundary", "A4 \u00b7 OT safety"),
    ("Appendix \u00b7 A5", "Assumptions register and 3-year P&L",
     ["Every input (A1\u2013A13) editable",
      "3-year scale path; improving ROI",
      "Structure holds; only inputs change"],
     "For finance deep-dives. Source: 06 \u00a70, \u00a74.",
     "A5", "06 \u00a70, \u00a74", "A5 \u00b7 P&L"),
    ("Appendix \u00b7 A6", "Compensation design",
     ["Influence credit is additive \u2014 BUs keep 100%",
      "50% bookings / 20% attach / 15% leading / 15% MBO",
      "Deal desk arbitrates via the ledger"],
     "For comp/quota questions. Source: 05 \u00a76.",
     "A6", "05 \u00a76", "A6 \u00b7 Comp"),
    ("Appendix \u00b7 A7", "Competitive framing",
     ["Point vendors, SIEM, SIs, hyperscalers",
      "SI dynamic is the biggest structural risk",
      "Whoever frames the architecture wins"],
     "For competitive/positioning questions. Source: 02 \u00a77.",
     "A7", "02 \u00a77", "A7 \u00b7 Competitive"),
    ("Appendix \u00b7 A8", "Full risk register + early warnings",
     ["12 risks: likelihood, impact, mitigation, owner",
      "Leading signals with action triggers",
      "Bounded downside via the Q3 gate"],
     "For risk/governance questions. Source: 07 \u00a73.",
     "A8", "07 \u00a73", "A8 \u00b7 Risk register"),
]

# Appendix divider with a contents grid
s = new_slide()
shp(s, MSO_SHAPE.RECTANGLE, 0.6, 0.9, 0.14, 1.35, ACCENT)
add_text(s, Inches(0.95), Inches(0.95), Inches(11.7), Inches(0.4), [[("APPENDIX", 13, ACCENT, True)]])
add_text(s, Inches(0.95), Inches(1.35), Inches(11.7), Inches(0.9),
         [[("Detail held in reserve for Q&A", 30, TEXT, True)]])
_cols, _cw2, _ch2 = 4, None, 0.85
_gx, _gy = 0.8, 3.1
_cw2 = (11.73 - 3 * 0.3) / 4
for _i, _row in enumerate(APPX):
    _lbl = _row[6]
    _x = _gx + (_i % _cols) * (_cw2 + 0.3)
    _y = _gy + (_i // _cols) * (_ch2 + 0.35)
    stext(shp(s, MSO_SHAPE.ROUNDED_RECTANGLE, _x, _y, _cw2, _ch2, CHIP, MUTED), _lbl, 12, TEXT, False)
s.notes_slide.notes_text_frame.text = "Use these only if the room asks. Don't present them by default."

for kicker, headline, bullets, notes, badge, srclabel, _toc in APPX:
    content_slide(kicker, headline, bullets, notes, visual=v_appendix(badge, srclabel),
                  footer="Appendix \u2014 reserve for Q&A. Illustrative figures.")

# ---- Page numbers (skip title slide) --------------------------------------
_total = len(prs.slides._sldIdLst)
for _i, _s in enumerate(prs.slides):
    if _i == 0:
        continue
    add_text(_s, Inches(11.5), Inches(7.03), Inches(1.23), Inches(0.3),
             [[("%d / %d" % (_i + 1, _total), 10, MUTED, False)]], align=PP_ALIGN.RIGHT)

prs.save("Field-CTO-XArch-Business-Case.pptx")
print("Wrote Field-CTO-XArch-Business-Case.pptx with", len(prs.slides._sldIdLst), "slides")
